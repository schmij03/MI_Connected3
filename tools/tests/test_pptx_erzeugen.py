"""Erzeugt eine Test-PowerPoint, die typische Folien der Lektionen nachbildet.

Wird von tools/tests/test_konverter.py verwendet. Inhalt frei erfunden.
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches, Pt

DUNKELBLAU = RGBColor(0x1F, 0x38, 0x64)
PETROL = RGBColor(0x1B, 0x6B, 0x73)
ORANGE = RGBColor(0xED, 0x7D, 0x31)
HELLGRAU = RGBColor(0xF2, 0xF2, 0xF2)
WEISS = RGBColor(0xFF, 0xFF, 0xFF)


def textbox(folie, x, y, w, h, zeilen, groesse=20, farbe=None, fett=False, fuellung=None, form=MSO_SHAPE.RECTANGLE):
    if fuellung is not None:
        s = folie.shapes.add_shape(form, Inches(x), Inches(y), Inches(w), Inches(h))
        s.fill.solid()
        s.fill.fore_color.rgb = fuellung
        s.line.fill.background()
    else:
        s = folie.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = s.text_frame
    tf.word_wrap = True
    for i, zeile in enumerate(zeilen):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        teile = zeile if isinstance(zeile, list) else [(zeile, fett)]
        for text, f in teile:
            r = p.add_run()
            r.text = text
            r.font.size = Pt(groesse)
            r.font.bold = f
            if farbe is not None:
                r.font.color.rgb = farbe
    return s


def titel(folie, text):
    return textbox(folie, 0.6, 0.35, 10, 0.9, [text], groesse=36, farbe=PETROL, fett=True)


def bild_png() -> io.BytesIO:
    im = Image.new("RGB", (640, 360), (242, 242, 242))
    d = ImageDraw.Draw(im)
    d.rectangle([40, 120, 200, 240], outline=(31, 56, 100), width=8)
    d.rectangle([260, 140, 380, 220], fill=(237, 125, 49))
    d.rectangle([440, 90, 600, 270], outline=(27, 107, 115), width=8)
    d.line([200, 180, 260, 180], fill=(27, 107, 115), width=6)
    d.line([380, 180, 440, 180], fill=(27, 107, 115), width=6)
    puffer = io.BytesIO()
    im.save(puffer, "PNG")
    puffer.seek(0)
    return puffer


def notiz(folie, text):
    folie.notes_slide.notes_text_frame.text = text


def erzeugen(ziel: Path) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    leer = prs.slide_layouts[6]
    titel_layout = prs.slide_layouts[1]

    # 1 Titelfolie: dunkelblau, Nummer im Kreis
    f = prs.slides.add_slide(leer)
    f.background.fill.solid()
    f.background.fill.fore_color.rgb = DUNKELBLAU
    textbox(f, 0.8, 2.4, 1.8, 1.8, ["2"], groesse=60, farbe=WEISS, fett=True, fuellung=ORANGE, form=MSO_SHAPE.OVAL)
    textbox(f, 3.0, 2.3, 9.5, 1.2, ["Kann man Bewertungen trauen?"], groesse=44, farbe=WEISS, fett=True)
    textbox(f, 3.0, 3.5, 9.5, 0.8, ["Und: Warum ruckelt das Video in den Ferien?"], groesse=26, farbe=WEISS)
    notiz(f, "Begrüssung, Ziele der Lektion an die Wandtafel schreiben.")

    # 2 Ausgeblendete Folie
    f = prs.slides.add_slide(leer)
    titel(f, "Planung (nur für mich)")
    textbox(f, 0.6, 1.5, 11, 3, ["Einstieg 5 Minuten", "Auftrag 1: 15 Minuten", "Besprechung 10 Minuten"], groesse=22)
    f._element.set("show", "0")

    # 3 Liste mit Platzhalter (Aufzaehlung vom Layout)
    f = prs.slides.add_slide(titel_layout)
    f.shapes.title.text = "Worum geht es heute?"
    body = f.placeholders[1].text_frame
    body.text = "Online-Bewertungen lesen und beurteilen"
    p = body.add_paragraph()
    p.text = "Gefälschte Bewertungen erkennen"
    p = body.add_paragraph()
    p.text = "Woran man sie erkennt"
    p.level = 1
    p = body.add_paragraph()
    r = p.add_run()
    r.text = "Warum "
    r = p.add_run()
    r.text = "Datenstau"
    r.font.bold = True
    r = p.add_run()
    r.text = " entsteht"
    notiz(f, "Kurz Vorwissen abfragen: Wer hat schon einmal etwas wegen einer Bewertung gekauft?")

    # 4 Auftragsfolie mit Seitenleiste und Zeitbadge
    f = prs.slides.add_slide(leer)
    textbox(
        f, 0, 0, 3.6, 7.5,
        ["AUFTRAG 1", "Zeit", "15'", "Sozialform", "Partnerarbeit", "Material", "Arbeitsblatt, Laptop", "Arbeitsbuch", "S. 12–13"],
        groesse=20, farbe=WEISS, fuellung=DUNKELBLAU,
    )
    textbox(f, 4.0, 0.4, 8, 0.9, ["Auftrag: Bewertungen prüfen"], groesse=34, farbe=PETROL, fett=True)
    tb = textbox(f, 4.0, 1.5, 8.5, 3.5, ["Lest die drei Bewertungen.", "Markiert verdächtige Stellen.", "Entscheidet, welcher ihr vertraut."], groesse=22)
    for p in tb.text_frame.paragraphs:
        pPr = p._p.get_or_add_pPr()
        from lxml import etree
        auto = etree.SubElement(pPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}buAutoNum")
        auto.set("type", "arabicPeriod")
    textbox(f, 11.8, 0.3, 1.2, 0.6, ["15'"], groesse=20, farbe=WEISS, fett=True, fuellung=ORANGE, form=MSO_SHAPE.ROUNDED_RECTANGLE)
    textbox(f, 4.0, 6.2, 8.8, 0.8, ["Tipp: Achtet auf Datum und Sprache der Bewertung."], groesse=18, fuellung=HELLGRAU)
    notiz(f, "Paare nach Sitzordnung.\nNach 10 Minuten Zwischenstand abfragen.")

    # 5 Tabelle
    f = prs.slides.add_slide(leer)
    titel(f, "Router, Provider, Server")
    tabelle = f.shapes.add_table(4, 3, Inches(0.6), Inches(1.6), Inches(12), Inches(3)).table
    daten = [["Begriff", "Was macht es?", "Beispiel"], ["Router", "verbindet Heimnetz und Internet", "WLAN-Router"],
             ["Provider", "verkauft den Internetzugang", "Swisscom | Salt"], ["Server", "liefert Daten", "YouTube"]]
    for z, zeile in enumerate(daten):
        for s, text in enumerate(zeile):
            tabelle.cell(z, s).text = text

    # 6 Text links, Bild rechts
    f = prs.slides.add_slide(leer)
    titel(f, "So reisen Daten")
    textbox(f, 0.6, 1.5, 6, 3, ["Daten werden in Pakete zerlegt.", "Jedes Paket findet seinen Weg."], groesse=22)
    bild = f.shapes.add_picture(bild_png(), Inches(7.2), Inches(1.5), Inches(5.5))
    bild._element.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr").set("descr", "Weg der Daten")

    # 7 Karten
    f = prs.slides.add_slide(leer)
    titel(f, "Drei Begriffe")
    for i, (kopf, text) in enumerate([("Router", "verteilt Daten im Heimnetz"), ("Provider", "Zugang zum Internet"), ("Server", "speichert Webseiten")]):
        textbox(f, 0.6 + i * 4.1, 1.8, 3.8, 3.2, [(kopf), text], groesse=22, fuellung=HELLGRAU)
        f.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True

    # 8 Loesung
    f = prs.slides.add_slide(leer)
    titel(f, "Lösung Auftrag 1")
    textbox(f, 0.6, 1.5, 12, 3, ["Bewertung A: eher nicht vertrauenswürdig", "Bewertung B: vertrauenswürdig"], groesse=22)

    # 9 Merke (dunkel)
    f = prs.slides.add_slide(leer)
    f.background.fill.solid()
    f.background.fill.fore_color.rgb = DUNKELBLAU
    textbox(f, 1.5, 2.5, 10, 2, ["Merke: Viele gleichzeitige Nutzer teilen sich die Leitung. Dann wird es langsam."], groesse=30, farbe=WEISS)

    ziel.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(ziel))
    return ziel


if __name__ == "__main__":
    print(erzeugen(Path(sys.argv[1])))
